from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.sales_analysis import (
    DEFAULT_INPUT_FILE,
    build_analysis_tables,
    build_kpi_summary,
    clean_sales_data,
    generate_outputs,
    load_and_clean_data,
)
from src.sales_analysis import _parse_date


def make_sales_rows() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["ORD-001", "2025-01-10", "Alice Jones", "Laptop", "Electronics", "North", 2, 1000, 200],
            [" ord-001 ", "2025-01-10", " alice  jones ", "laptop", "electronics", "north", 2, "1,000", "200"],
            ["ORD-002", "2025-13-10", "Alice Jones", "Laptop", "Electronics", "North", -1, -20, -4],
            ["ORD-003", None, None, "Monitor", "Electronics", "East", None, None, None],
            ["ORD-004", "2025-02-10", "Bob Smith", "Desk Lamp", "Furniture", "South", 1, 50, 12],
        ],
        columns=[
            "Order ID",
            "Order Date",
            "Customer",
            "Product",
            "Category",
            "Region",
            "Quantity",
            "Sales",
            "Profit",
        ],
    )


class CleanSalesDataTests(unittest.TestCase):
    def test_cleans_normalized_duplicates_without_modifying_source(self) -> None:
        raw = make_sales_rows()
        original = raw.copy(deep=True)

        cleaned, report = clean_sales_data(raw)

        self.assertEqual(raw.shape, original.shape)
        pd.testing.assert_frame_equal(raw, original)
        self.assertEqual(len(cleaned), 4)
        self.assertEqual(cleaned["Order ID"].tolist(), ["ORD-001", "ORD-002", "ORD-003", "ORD-004"])
        self.assertEqual(cleaned.loc[0, "Customer"], "Alice Jones")
        self.assertEqual(cleaned.loc[0, "Category"], "Electronics")
        self.assertEqual(cleaned.loc[1, "Sales"], -20.0)
        self.assertEqual(cleaned.loc[2, "Customer"], "Unknown")
        metrics = dict(zip(report["Metric"], report["Value"]))
        self.assertEqual(metrics["Duplicate rows removed"], 1)
        self.assertEqual(metrics["Rows with missing or invalid order dates"], 2)
        self.assertEqual(metrics["Rows with negative sales"], 1)
        self.assertEqual(metrics["Rows with zero or negative quantity"], 1)
        self.assertEqual(metrics["Rows with missing or invalid sales"], 1)
        self.assertEqual(metrics["Rows with missing or invalid profit"], 1)

    def test_keeps_distinct_line_items_with_the_same_order_id(self) -> None:
        raw = make_sales_rows().iloc[[0, 0]].copy()
        raw.iloc[1, raw.columns.get_loc("Product")] = "Monitor"

        cleaned, report = clean_sales_data(raw)

        self.assertEqual(len(cleaned), 2)
        self.assertEqual(cleaned["Order ID"].nunique(), 1)
        self.assertEqual(dict(zip(report["Metric"], report["Value"]))["Duplicate rows removed"], 0)

    def test_invalid_input_schema_raises_clear_error(self) -> None:
        with self.assertRaisesRegex(ValueError, "Missing required sales columns: Order ID"):
            clean_sales_data(pd.DataFrame({"Order Date": ["2025-01-01"]}))

    def test_empty_input_raises_clear_error(self) -> None:
        with self.assertRaisesRegex(ValueError, "has no rows"):
            clean_sales_data(pd.DataFrame(columns=["Order ID"]))

    def test_excel_serial_dates_are_parsed_correctly(self) -> None:
        self.assertEqual(
            _parse_date(45728),
            pd.Timestamp("2025-03-12"),
        )

    def test_rows_without_order_ids_are_reported_and_excluded(self) -> None:
        raw = make_sales_rows().iloc[[0]].copy()
        raw.loc[raw.index[0], "Order ID"] = "  "
        cleaned, report = clean_sales_data(raw)

        self.assertTrue(cleaned.empty)
        self.assertEqual(
            dict(zip(report["Metric"], report["Value"]))[
                "Rows removed without an order ID"
            ],
            1,
        )

    def test_kpis_exclude_unknown_customers_and_keep_unknown_sales_unknown(self) -> None:
        cleaned, quality = clean_sales_data(make_sales_rows())

        kpis = build_kpi_summary(cleaned, quality).set_index("Metric")["Value"]

        self.assertEqual(kpis["Total Sales"], 1030.0)
        self.assertEqual(kpis["Total Profit"], 208.0)
        self.assertEqual(kpis["Total Orders"], 4)
        self.assertEqual(kpis["Total Customers"], 2)
        self.assertEqual(kpis["Average Order Value"], 257.5)
        self.assertEqual(kpis["Profit Margin"], 208 / 1030)

    def test_monthly_trends_exclude_invalid_dates_but_preserve_negative_sales(self) -> None:
        cleaned, _ = clean_sales_data(make_sales_rows())
        tables = build_analysis_tables(cleaned)

        monthly = tables["monthly"]
        self.assertEqual(len(monthly), 2)
        self.assertEqual(float(monthly.loc[0, "Sales"]), 1000.0)
        self.assertTrue(pd.isna(monthly.loc[0, "Sales Growth %"]))
        self.assertEqual(float(monthly.loc[1, "Sales"]), 50.0)
        self.assertEqual(tables["customers"].loc[0, "Customer"], "Alice Jones")
        self.assertEqual(int(tables["customers"].loc[0, "Orders"]), 2)

    def test_all_invalid_dates_produce_empty_time_series_without_losing_sales(self) -> None:
        raw = make_sales_rows().iloc[[1]].copy()
        raw["Order Date"] = "not a date"

        cleaned, _ = clean_sales_data(raw)
        tables = build_analysis_tables(cleaned)
        kpis = build_kpi_summary(cleaned).set_index("Metric")["Value"]

        self.assertEqual(len(tables["monthly"]), 0)
        self.assertEqual(kpis["Total Sales"], 1000.0)

    def test_supplied_workbook_has_cleanable_real_source_rows(self) -> None:
        cleaned, report = load_and_clean_data(DEFAULT_INPUT_FILE)
        metrics = dict(zip(report["Metric"], report["Value"]))

        self.assertEqual(metrics["Source data rows"], 1500)
        self.assertEqual(metrics["Exact duplicate rows in source"], 2)
        self.assertEqual(metrics["Rows with negative sales"], 1)
        self.assertEqual(metrics["Rows with zero or negative quantity"], 2)
        self.assertGreater(metrics["Rows with missing or invalid order dates"], 0)
        self.assertGreater(len(cleaned), 1000)
        self.assertTrue(set(("Order ID", "Order Date", "Sales", "Profit")).issubset(cleaned.columns))


class OutputGenerationTests(unittest.TestCase):
    def test_generates_excel_report_powerbi_csv_files_and_visualizations(self) -> None:
        cleaned, quality = clean_sales_data(make_sales_rows())

        with tempfile.TemporaryDirectory() as temporary:
            output_dir = Path(temporary)
            outputs = generate_outputs(
                cleaned,
                quality,
                output_dir=output_dir,
            )

            self.assertTrue(outputs["cleaned_sales"].is_file())
            self.assertTrue(outputs["kpi_summary"].is_file())
            self.assertTrue(outputs["data_quality_report"].is_file())
            self.assertTrue(outputs["excel_analysis"].is_file())
            self.assertTrue(outputs["business_insights"].is_file())
            self.assertTrue((output_dir / "sales_by_month.csv").is_file())
            self.assertTrue((output_dir / "charts" / "monthly_sales.png").is_file())
            self.assertTrue(
                (output_dir / "charts" / "sales_by_category.png").is_file()
            )
            self.assertTrue(
                (output_dir / "charts" / "top_10_products.png").is_file()
            )
            self.assertEqual(len(pd.read_csv(outputs["cleaned_sales"])), 4)

            workbook = load_workbook(outputs["excel_analysis"], read_only=False)
            self.assertEqual(workbook.sheetnames, ["Dashboard"])
            dashboard = workbook["Dashboard"]
            self.assertEqual(len(dashboard._charts), 5)
            self.assertIn("DashMonthly", dashboard.tables)
            self.assertIn("DashCustomers", dashboard.tables)
            self.assertIn("DashQuality", dashboard.tables)
            self.assertIn("DashCleanedSales", dashboard.tables)
            cleaned_table = dashboard.tables["DashCleanedSales"]
            first_data_row = int(cleaned_table.ref.split(":")[0][1:]) + 1
            self.assertEqual(
                dashboard.cell(row=first_data_row, column=1).value,
                "ORD-001",
            )
            self.assertGreater(dashboard.max_row, 10)
            workbook.close()

    def test_generated_insights_disclose_unspecified_currency(self) -> None:
        cleaned, quality = clean_sales_data(make_sales_rows())

        with tempfile.TemporaryDirectory() as temporary:
            outputs = generate_outputs(cleaned, quality, output_dir=temporary)
            insights = outputs["business_insights"].read_text(encoding="utf-8")

        self.assertIn("source does not specify a currency", insights)
        self.assertIn("Highest-sales product", insights)
        self.assertIn("negative sales", insights)


if __name__ == "__main__":
    unittest.main()
