"""Clean, analyze, and export the sales internship project."""

from __future__ import annotations

import argparse
import math
from numbers import Real
from datetime import date, datetime
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT_FILE = PROJECT_ROOT / "data" / "raw" / "Sales_Data_Raw_Practice.xlsx"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "outputs"
EXPECTED_COLUMNS = (
    "Order ID",
    "Order Date",
    "Customer",
    "Product",
    "Category",
    "Region",
    "Quantity",
    "Sales",
    "Profit",
)
TEXT_LABEL_COLUMNS = ("Customer", "Product", "Category", "Region")
NUMERIC_COLUMNS = ("Quantity", "Sales", "Profit")
CHART_COLOR = "#187B86"
DARK_COLOR = "#17324D"


def _missing_mask(series: pd.Series) -> pd.Series:
    """Identify actual missing values as well as blank and whitespace-only text."""
    return series.isna() | series.astype("string").str.strip().eq("").fillna(False)


def _parse_number(value: Any) -> Any:
    """Parse ordinary and thousands-separated numbers without inventing values."""
    if pd.isna(value):
        return pd.NA
    cleaned = str(value).strip().replace(",", "")
    if cleaned.startswith("(") and cleaned.endswith(")"):
        cleaned = f"-{cleaned[1:-1]}"
    parsed = pd.to_numeric(cleaned, errors="coerce")
    return pd.NA if pd.isna(parsed) else float(parsed)


def _parse_date(value: Any) -> Any:
    """Parse source dates, including Excel serials, while retaining invalid values as NaT."""
    if pd.isna(value):
        return pd.NaT
    if isinstance(value, Real) and not isinstance(value, bool):
        parsed = pd.to_datetime(
            value, unit="D", origin="1899-12-30", errors="coerce", utc=True
        )
    else:
        parsed = pd.to_datetime(value, format="mixed", errors="coerce", utc=True)
    if pd.isna(parsed):
        return pd.NaT
    return parsed.tz_localize(None)


def _numeric_or_missing_count(series: pd.Series) -> int:
    """Count nonblank source values that cannot be parsed as a number."""
    source_missing = _missing_mask(series)
    parsed = series.map(_parse_number)
    return int((parsed.isna() & ~source_missing).sum())


def clean_sales_data(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return a cleaned copy and a transparent, row-counted data-quality report."""
    if raw.empty:
        raise ValueError("The supplied sales dataset has no rows.")

    cleaned_headers = [str(header).strip() for header in raw.columns]
    if len(cleaned_headers) != len(set(cleaned_headers)):
        raise ValueError("The dataset has duplicate column names after trimming spaces.")
    missing_columns = [column for column in EXPECTED_COLUMNS if column not in cleaned_headers]
    if missing_columns:
        raise ValueError(
            "Missing required sales columns: " + ", ".join(missing_columns)
        )

    source = raw.copy()
    source.columns = cleaned_headers
    source_rows = len(source)
    source_exact_duplicates = int(source.duplicated(keep="first").sum())
    missing_by_column = {
        column: int(_missing_mask(source[column]).sum()) for column in EXPECTED_COLUMNS
    }

    working = source.copy()
    for column in ("Order ID", *TEXT_LABEL_COLUMNS):
        text = (
            working[column]
            .astype("string")
            .str.replace(r"\s+", " ", regex=True)
            .str.strip()
            .replace("", pd.NA)
        )
        if column == "Order ID":
            working[column] = text.str.upper()
        else:
            working[column] = text.str.title().fillna("Unknown")

    invalid_numeric_counts = {
        column: _numeric_or_missing_count(source[column])
        for column in NUMERIC_COLUMNS
    }
    for column in NUMERIC_COLUMNS:
        working[column] = working[column].map(_parse_number).astype("Float64")

    working["Order Date"] = working["Order Date"].map(_parse_date)
    duplicate_rows_removed = int(working.duplicated(keep="first").sum())
    working = working.drop_duplicates(keep="first").copy()
    missing_order_ids = _missing_mask(working["Order ID"])
    rows_without_order_id = int(missing_order_ids.sum())
    working = working.loc[~missing_order_ids].reset_index(drop=True)
    invalid_dates = int(working["Order Date"].isna().sum())
    negative_sales = int((working["Sales"].dropna() < 0).sum())
    nonpositive_quantity = int((working["Quantity"].dropna() <= 0).sum())
    missing_sales = int(working["Sales"].isna().sum())
    missing_profit = int(working["Profit"].isna().sum())
    missing_quantity = int(working["Quantity"].isna().sum())

    report_rows: list[dict[str, Any]] = [
        {
            "Metric": "Source data rows",
            "Value": source_rows,
            "Details": "Number of records in the untouched source worksheet.",
        },
        {
            "Metric": "Exact duplicate rows in source",
            "Value": source_exact_duplicates,
            "Details": "Rows identical across all original columns.",
        },
        {
            "Metric": "Duplicate rows removed",
            "Value": duplicate_rows_removed,
            "Details": "Identical rows removed after normalizing text and numeric values. Distinct line items with the same Order ID are retained.",
        },
        {
            "Metric": "Rows removed without an order ID",
            "Value": rows_without_order_id,
            "Details": "An order ID is required to count orders without counting unknown IDs as real orders.",
        },
        {
            "Metric": "Cleaned data rows",
            "Value": len(working),
            "Details": "Rows remaining after documented deduplication and removal of records without an order ID.",
        },
        {
            "Metric": "Rows with missing or invalid order dates",
            "Value": invalid_dates,
            "Details": "Retained in sales totals; excluded from time-series charts and date-based trend comparisons.",
        },
        {
            "Metric": "Rows with missing or invalid sales",
            "Value": missing_sales,
            "Details": "Retained with an unknown Sales value; sales totals skip unknown amounts rather than treating them as zero.",
        },
        {
            "Metric": "Rows with missing or invalid profit",
            "Value": missing_profit,
            "Details": "Retained with an unknown Profit value; profit totals skip unknown amounts rather than treating them as zero.",
        },
        {
            "Metric": "Rows with missing or invalid quantity",
            "Value": missing_quantity,
            "Details": "Retained with an unknown Quantity; units-sold totals include only reported quantities.",
        },
        {
            "Metric": "Rows with negative sales",
            "Value": negative_sales,
            "Details": "Retained unchanged; confirm whether negative sales represent returns, credits, or source errors.",
        },
        {
            "Metric": "Rows with zero or negative quantity",
            "Value": nonpositive_quantity,
            "Details": "Retained unchanged; confirm whether zero or negative quantities are intentional returns or source errors.",
        },
    ]
    for column in EXPECTED_COLUMNS:
        report_rows.append(
            {
                "Metric": f"Source rows missing {column}",
                "Value": missing_by_column[column],
                "Details": "Count in the original worksheet before cleaning.",
            }
        )
    for column, count in invalid_numeric_counts.items():
        if count:
            report_rows.append(
                {
                    "Metric": f"Unparseable {column} values",
                    "Value": count,
                    "Details": "Nonblank source values retained as unknown instead of being guessed.",
                }
            )

    return working, pd.DataFrame(report_rows, columns=["Metric", "Value", "Details"])


def load_and_clean_data(
    input_file: str | Path = DEFAULT_INPUT_FILE,
    sheet_name: str = "Raw_Sales_Data",
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Load an Excel worksheet or CSV file, then clean a new working copy."""
    path = Path(input_file)
    if not path.is_file():
        raise FileNotFoundError(f"Sales dataset does not exist: {path}")
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        raw = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")
    elif suffix == ".csv":
        raw = pd.read_csv(path, dtype={"Order ID": "string"})
    else:
        raise ValueError(
            f"Unsupported dataset format {suffix!r}. Please provide a CSV or Excel workbook."
        )
    return clean_sales_data(raw)


def _sum(values: pd.Series) -> Any:
    total = values.sum(min_count=1)
    return pd.NA if pd.isna(total) else float(total)


def _profit_margin(frame: pd.DataFrame) -> pd.Series:
    return (frame["Profit"] / frame["Sales"]).where(frame["Sales"].ne(0))


def build_kpi_summary(
    sales: pd.DataFrame, quality_report: pd.DataFrame | None = None
) -> pd.DataFrame:
    """Calculate consistent, clearly defined project KPIs from cleaned sales rows."""
    total_sales = _sum(sales["Sales"])
    total_profit = _sum(sales["Profit"])
    total_orders = int(sales["Order ID"].nunique())
    units_sold = _sum(sales["Quantity"])
    known_customers = int(sales.loc[sales["Customer"].ne("Unknown"), "Customer"].nunique())
    average_order_value = (
        total_sales / total_orders
        if total_sales is not pd.NA and total_orders
        else pd.NA
    )
    profit_margin = (
        total_profit / total_sales
        if total_sales is not pd.NA and total_sales != 0 and total_profit is not pd.NA
        else pd.NA
    )

    report_lookup = (
        dict(zip(quality_report["Metric"], quality_report["Value"]))
        if quality_report is not None
        else {}
    )
    metrics = [
        ("Total Sales", total_sales, "Sum of known sales; source currency is unspecified."),
        ("Total Profit", total_profit, "Sum of known profit values; missing profit is not treated as zero."),
        ("Total Orders", total_orders, "Distinct, nonblank Order IDs; repeated IDs can represent valid line items."),
        ("Units Sold", units_sold, "Sum of known quantities; review the data-quality report for missing or nonpositive quantities."),
        ("Total Customers", known_customers, "Distinct customers with an identified name; Unknown is excluded."),
        ("Average Order Value", average_order_value, "Total known sales divided by distinct orders."),
        ("Profit Margin", profit_margin, "Total known profit divided by total known sales; undefined when sales are zero or profit is unknown."),
        ("Rows with invalid order dates", report_lookup.get("Rows with missing or invalid order dates", 0), "Retained in total sales but excluded from time-series analysis."),
        ("Rows with missing or invalid sales", report_lookup.get("Rows with missing or invalid sales", 0), "Missing amounts remain unknown, not zero."),
        ("Rows with missing or invalid profit", report_lookup.get("Rows with missing or invalid profit", 0), "Missing amounts remain unknown, not zero."),
        ("Duplicate rows removed", report_lookup.get("Duplicate rows removed", 0), "Normalized identical rows removed; distinct order line items retained."),
    ]
    return pd.DataFrame(metrics, columns=["Metric", "Value", "Definition"])


def _aggregate_by_dimension(sales: pd.DataFrame, column: str) -> pd.DataFrame:
    grouped = sales.groupby(column, dropna=False, observed=True)
    result = grouped.agg(
        Sales=("Sales", _sum),
        Profit=("Profit", _sum),
        **{"Units Sold": ("Quantity", _sum)},
        Orders=("Order ID", "nunique"),
    ).reset_index()
    result["Profit Margin"] = _profit_margin(result)
    return result.sort_values(
        "Sales", ascending=False, na_position="last", kind="stable"
    ).reset_index(drop=True)


def _grouped_period_summary(
    dated: pd.DataFrame, group_column: str, label: str
) -> pd.DataFrame:
    grouped = dated.groupby(group_column, observed=True).agg(
        Sales=("Sales", _sum),
        Profit=("Profit", _sum),
        Orders=("Order ID", "nunique"),
        **{"Units Sold": ("Quantity", _sum)},
    )
    grouped.index.name = label
    result = grouped.reset_index()
    result["Sales Growth %"] = result["Sales"].pct_change(fill_method=None).replace(
        [float("inf"), float("-inf")], pd.NA
    )
    return result


def build_analysis_tables(sales: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Prepare the period, product, category, regional, and customer analyses."""
    dated = sales.loc[sales["Order Date"].notna()].copy()
    tables: dict[str, pd.DataFrame] = {}

    if dated.empty:
        tables["monthly"] = pd.DataFrame(
            columns=["Month", "Sales", "Profit", "Orders", "Units Sold", "Sales Growth %"]
        )
        tables["quarterly"] = pd.DataFrame(
            columns=["Quarter", "Sales", "Profit", "Orders", "Units Sold", "Sales Growth %"]
        )
        tables["yearly"] = pd.DataFrame(
            columns=["Year", "Sales", "Profit", "Orders", "Units Sold", "Sales Growth %"]
        )
    else:
        dated["_Month"] = dated["Order Date"].dt.to_period("M").dt.to_timestamp()
        monthly = _grouped_period_summary(dated, "_Month", "Month")
        month_range = pd.date_range(
            monthly["Month"].min(), monthly["Month"].max(), freq="MS", name="Month"
        )
        monthly = monthly.set_index("Month").reindex(month_range)
        empty_months = ~monthly.index.isin(
            dated["_Month"].drop_duplicates().to_numpy()
        )
        for column in ("Sales", "Profit", "Orders", "Units Sold"):
            monthly.loc[empty_months, column] = (
                0 if column == "Orders" else (0.0 if column in ("Sales", "Profit") else pd.NA)
            )
        monthly["Sales Growth %"] = monthly["Sales"].pct_change(fill_method=None).replace(
            [float("inf"), float("-inf")], pd.NA
        )
        tables["monthly"] = monthly.rename_axis("Month").reset_index()

        dated["_Quarter"] = dated["Order Date"].dt.to_period("Q").astype("string")
        tables["quarterly"] = _grouped_period_summary(dated, "_Quarter", "Quarter")
        dated["_Year"] = dated["Order Date"].dt.year
        tables["yearly"] = _grouped_period_summary(dated, "_Year", "Year")

    tables["products"] = _aggregate_by_dimension(sales, "Product")
    tables["categories"] = _aggregate_by_dimension(sales, "Category")
    tables["regions"] = _aggregate_by_dimension(sales, "Region")
    customers = sales.loc[sales["Customer"].ne("Unknown")].copy()
    tables["customers"] = _aggregate_by_dimension(customers, "Customer")
    tables["customers"]["Repeat Customer"] = tables["customers"]["Orders"].gt(1)
    return tables


def _kpi_value(kpis: pd.DataFrame, name: str) -> Any:
    match = kpis.loc[kpis["Metric"].eq(name), "Value"]
    return match.iloc[0] if not match.empty else pd.NA


def build_insights(
    sales: pd.DataFrame,
    quality_report: pd.DataFrame,
    tables: dict[str, pd.DataFrame],
    kpis: pd.DataFrame,
) -> str:
    """Create cautious, data-derived insights and practical follow-up recommendations."""
    total_sales = _kpi_value(kpis, "Total Sales")
    lines = [
        "# Sales analysis: findings and recommendations",
        "",
        "> Source amounts are reported in the dataset's units. The source does not specify a currency.",
        "",
        "## Data and methodology",
        "",
        f"- Cleaned records analyzed: **{len(sales):,}**.",
        "- Order counts use distinct Order IDs; valid multiple line items are preserved.",
        "- Identical records after text normalization are removed, but suspicious returns or values are not automatically deleted.",
        "- Rows with missing or invalid dates contribute to available sales totals but are excluded from monthly, quarterly, and yearly analyses.",
        "- Missing Sales, Profit, and Quantity values are kept unknown rather than filled with zero.",
        "- Profit margin is calculated from total known profit divided by total known sales.",
        "",
        "## Findings",
        "",
    ]

    def add_leader(table_name: str, field: str, value_column: str, description: str) -> None:
        table = tables[table_name]
        valid = table.loc[table[value_column].notna()]
        if valid.empty:
            lines.append(f"- **{description}:** No complete {value_column.lower()} data is available.")
            return
        leader = valid.iloc[0]
        value = float(leader[value_column])
        share = ""
        if (
            value_column == "Sales"
            and total_sales is not pd.NA
            and float(total_sales) != 0
        ):
            share = f" ({value / float(total_sales):.1%} of net sales)"
        lines.append(
            f"- **{description}:** {leader[field]} - {value:,.2f}{share}."
        )

    add_leader("products", "Product", "Sales", "Highest-sales product")
    add_leader("categories", "Category", "Sales", "Highest-sales category")
    add_leader("regions", "Region", "Sales", "Highest-sales region")

    monthly = tables["monthly"]
    valid_months = monthly.loc[monthly["Sales"].notna()]
    if not valid_months.empty:
        best_month = valid_months.loc[valid_months["Sales"].idxmax()]
        lines.append(
            f"- **Highest-sales month:** {best_month['Month']:%B %Y} — "
            f"{float(best_month['Sales']):,.2f}."
        )
        growth = valid_months["Sales Growth %"].dropna()
        if not growth.empty:
            latest_growth = float(growth.iloc[-1])
            lines.append(
                f"- **Change in the latest comparable month:** {latest_growth:+.1%}. "
                "This compares with the immediately preceding calendar month, not the same month last year."
            )
    else:
        lines.append("- **Sales trend:** There are no valid transaction dates for a time-series comparison.")

    customers = tables["customers"]
    if not customers.empty:
        repeat_count = int(customers["Repeat Customer"].sum())
        lines.append(
            f"- **Repeat customers:** {repeat_count:,} of {len(customers):,} identified customers "
            f"({repeat_count / len(customers):.1%}) placed orders on more than one distinct Order ID."
        )
    else:
        lines.append("- **Customer retention:** No identified customer records are available.")

    quality = dict(zip(quality_report["Metric"], quality_report["Value"]))
    review_items = [
        (
            "Rows with missing or invalid order dates",
            "record has a missing or invalid date and is excluded from time-based charts",
            "records have missing or invalid dates and are excluded from time-based charts",
        ),
        (
            "Rows with negative sales",
            "record has negative sales; confirm whether it is a return or credit",
            "records have negative sales; confirm whether they are returns or credits",
        ),
        (
            "Rows with zero or negative quantity",
            "record has zero or negative quantity; confirm return conventions",
            "records have zero or negative quantities; confirm return conventions",
        ),
        (
            "Rows with missing or invalid profit",
            "record has missing or invalid profit",
            "records have missing or invalid profit",
        ),
    ]
    lines.extend(["", "## Data-quality issues to review", ""])
    found_review_item = False
    for key, singular_description, plural_description in review_items:
        count = int(quality.get(key, 0))
        if count:
            description = (
                singular_description if count == 1 else plural_description
            )
            lines.append(f"- **{count:,} {description}.**")
            found_review_item = True
    if not found_review_item:
        lines.append("- No missing-date, negative-sales, nonpositive-quantity, or missing-profit issues were identified.")

    lines.extend(["", "## Recommended next steps", ""])
    if not customers.empty and not customers["Repeat Customer"].any():
        lines.append(
            "- Review customer/order identifiers and the reporting period before setting retention targets; the available data shows no repeat identified customers."
        )
    else:
        lines.append(
            "- Compare repeat-customer revenue and order frequency with first-time customers before choosing retention incentives."
        )
    if not tables["products"].empty:
        top_product = tables["products"].loc[tables["products"]["Sales"].notna()]
        if not top_product.empty:
            lines.append(
                f"- Review availability, margins, and customer demand for **{top_product.iloc[0]['Product']}**, "
                "the highest-sales product; prioritize profitable growth rather than revenue alone."
            )
    if not tables["categories"].empty:
        profit_categories = tables["categories"].loc[tables["categories"]["Profit"].notna()]
        if not profit_categories.empty:
            best_profit = profit_categories.sort_values(
                "Profit", ascending=False, kind="stable"
            ).iloc[0]
            lines.append(
                f"- Check whether **{best_profit['Category']}**, the highest reported total-profit category, "
                "also performs well on sales and margin; protect profitable categories when adjusting the product mix."
            )
    lines.append(
        "- Ask the source owner to confirm currency, negative-quantity conventions, negative sales, and missing profit before finalizing business or financial decisions."
    )

    return "\n".join(lines) + "\n"


def _format_number(value: Any) -> str:
    if pd.isna(value):
        return "Not available"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    number = float(value)
    if number.is_integer():
        return f"{int(number):,}"
    return f"{number:,.2f}"


def _format_chart_amount(value: float, _position: float) -> str:
    if abs(value) >= 1_000_000:
        return f"{value / 1_000_000:,.1f}M"
    if abs(value) >= 1_000:
        return f"{value / 1_000:,.0f}K"
    return f"{value:,.0f}"


def _write_visualizations(
    tables: dict[str, pd.DataFrame], charts_dir: Path
) -> dict[str, Path]:
    charts_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    plt.rcParams.update(
        {
            "figure.facecolor": "#FFFFFF",
            "axes.facecolor": "#FFFFFF",
            "axes.edgecolor": "#DCE5EC",
            "axes.labelcolor": DARK_COLOR,
            "text.color": DARK_COLOR,
            "xtick.color": "#5C6B78",
            "ytick.color": "#5C6B78",
            "font.family": "DejaVu Sans",
        }
    )

    monthly = tables["monthly"]
    fig, ax = plt.subplots(figsize=(11, 5))
    if not monthly.empty and monthly["Sales"].notna().any():
        ax.plot(
            monthly["Month"],
            monthly["Sales"],
            color=CHART_COLOR,
            linewidth=2.5,
            marker="o",
            markersize=4,
        )
        ax.xaxis.set_major_locator(mdates.AutoDateLocator())
        ax.xaxis.set_major_formatter(mdates.ConciseDateFormatter(ax.xaxis.get_major_locator()))
        ax.grid(axis="y", color="#E8EDF1", linewidth=0.8)
        ax.set_ylabel("Sales (source units; currency unspecified)")
        ax.set_xlabel("")
        ax.yaxis.set_major_formatter(FuncFormatter(_format_chart_amount))
    else:
        ax.text(0.5, 0.5, "No valid order dates to plot.", ha="center", va="center", transform=ax.transAxes)
        ax.set_axis_off()
    ax.set_title("Monthly sales", loc="left", fontweight="bold", pad=16)
    fig.tight_layout()
    path = charts_dir / "monthly_sales.png"
    fig.savefig(path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    paths["monthly_sales"] = path

    visual_specs = (
        ("categories", "Category", "Sales by category", "sales_by_category.png"),
        ("products", "Product", "Top 10 products by sales", "top_10_products.png"),
        ("regions", "Region", "Sales by region", "sales_by_region.png"),
        ("customers", "Customer", "Top 10 customers by sales", "top_10_customers.png"),
    )
    for table_key, dimension, title, filename in visual_specs:
        table = tables[table_key].loc[tables[table_key]["Sales"].notna()].head(10)
        fig, ax = plt.subplots(figsize=(9, 5))
        if table.empty:
            ax.text(0.5, 0.5, "No reported sales to plot.", ha="center", va="center", transform=ax.transAxes)
            ax.set_axis_off()
        else:
            displayed = table.iloc[::-1]
            ax.barh(displayed[dimension].astype(str), displayed["Sales"], color=CHART_COLOR)
            ax.set_xlabel("Sales (source units; currency unspecified)")
            ax.grid(axis="x", color="#E8EDF1", linewidth=0.8)
            ax.set_axisbelow(True)
            ax.xaxis.set_major_formatter(FuncFormatter(_format_chart_amount))
        ax.set_title(title, loc="left", fontweight="bold", pad=16)
        fig.tight_layout()
        path = charts_dir / filename
        fig.savefig(path, dpi=160, bbox_inches="tight")
        plt.close(fig)
        paths[table_key] = path

    return paths


def _excel_value(value: Any) -> Any:
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, (datetime, date, str, int, float, bool)):
        return None if isinstance(value, float) and not math.isfinite(value) else value
    item = getattr(value, "item", None)
    if callable(item):
        return _excel_value(item())
    missing = pd.isna(value)
    return None if isinstance(missing, bool) and missing else str(value)


def _write_dashboard_table(
    sheet: Any,
    title: str,
    frame: pd.DataFrame,
    start_row: int,
    start_column: int,
    table_name: str,
) -> int:
    """Write a formatted, filterable analysis table into the single dashboard sheet."""
    headers = [str(column) for column in frame.columns]
    if not headers:
        return start_row

    end_column = start_column + len(headers) - 1
    title_cell = sheet.cell(row=start_row, column=start_column, value=title)
    title_cell.font = Font(bold=True, color="17324D", size=11)
    title_cell.fill = PatternFill("solid", fgColor="EAF1F5")
    if end_column > start_column:
        sheet.merge_cells(
            start_row=start_row,
            start_column=start_column,
            end_row=start_row,
            end_column=end_column,
        )
        for column in range(start_column, end_column + 1):
            sheet.cell(row=start_row, column=column).fill = PatternFill(
                "solid", fgColor="EAF1F5"
            )

    header_row = start_row + 1
    for offset, header in enumerate(headers):
        cell = sheet.cell(row=header_row, column=start_column + offset, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 28

    for row_offset, values in enumerate(
        frame.itertuples(index=False, name=None), start=header_row + 1
    ):
        for column_offset, value in enumerate(values):
            cell = sheet.cell(
                row=row_offset,
                column=start_column + column_offset,
                value=_excel_value(value),
            )
            header = headers[column_offset]
            if "Growth %" in header or header == "Profit Margin":
                cell.number_format = "0.0%;[Red]-0.0%;-"
            elif header in {"Order Date", "Month"}:
                cell.number_format = "dd mmm yyyy"
            elif header == "Year":
                cell.number_format = "0"
            elif isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00;[Red]-#,##0.00;-"

    last_row = header_row + len(frame)
    if len(frame):
        table = Table(
            displayName=table_name,
            ref=(
                f"{get_column_letter(start_column)}{header_row}:"
                f"{get_column_letter(end_column)}{last_row}"
            ),
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

        for offset, header in enumerate(headers):
            if header not in {"Sales", "Profit"}:
                continue
            column_letter = get_column_letter(start_column + offset)
            sheet.conditional_formatting.add(
                f"{column_letter}{header_row + 1}:{column_letter}{last_row}",
                ColorScaleRule(
                    start_type="min",
                    start_color="FCE8E6",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFF4D6",
                    end_type="max",
                    end_color="D5F0E3",
                ),
            )

    for offset, header in enumerate(headers):
        column = start_column + offset
        samples = [header] + [
            str(value)
            for value in frame.iloc[:50, offset].tolist()
            if value is not None and not pd.isna(value)
        ]
        estimated_width = min(27, max(12, max(map(len, samples)) + 2))
        letter = get_column_letter(column)
        current_width = sheet.column_dimensions[letter].width or 0
        sheet.column_dimensions[letter].width = min(
            30, max(current_width, estimated_width)
        )
    return last_row


def _add_dashboard_insights(sheet: Any, insights: str, start_row: int) -> None:
    sections: dict[str, list[str]] = {}
    section = ""
    for line in insights.splitlines():
        if line.startswith("## "):
            section = line[3:].strip()
            sections[section] = []
        elif section and line.startswith("- "):
            sections[section].append(
                line[2:].replace("**", "").replace("`", "").strip()
            )

    for start_column, end_column, title, key, fill in (
        (1, 11, "KEY FINDINGS", "Findings", "EAF1F5"),
        (12, 21, "RECOMMENDATIONS", "Recommended next steps", "F3F6F8"),
    ):
        sheet.merge_cells(
            start_row=start_row,
            start_column=start_column,
            end_row=start_row,
            end_column=end_column,
        )
        heading = sheet.cell(row=start_row, column=start_column, value=title)
        heading.font = Font(bold=True, color="17324D", size=11)
        for column in range(start_column, end_column + 1):
            sheet.cell(row=start_row, column=column).fill = PatternFill(
                "solid", fgColor=fill
            )
        for offset, message in enumerate(sections.get(key, [])[:5], start=1):
            row = start_row + offset
            sheet.merge_cells(
                start_row=row,
                start_column=start_column,
                end_row=row,
                end_column=end_column,
            )
            cell = sheet.cell(row=row, column=start_column, value=f"• {message}")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(color="34475A", size=10)
            sheet.row_dimensions[row].height = 32


def _add_dashboard_sheet(
    sheet: Any,
    kpis: pd.DataFrame,
    analysis_tables: dict[str, pd.DataFrame],
    quality_report: pd.DataFrame,
    sales: pd.DataFrame,
    insights: str,
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 70
    sheet.freeze_panes = "A9"
    sheet["A1"] = "SALES ANALYSIS | PROJECT DASHBOARD"
    sheet["A1"].font = Font(name="Aptos Display", size=20, bold=True, color="17324D")
    sheet.merge_cells("A1:N1")
    sheet["A2"] = "Sales amounts use the units reported in the source; no currency was specified."
    sheet["A2"].font = Font(italic=True, color="657587")
    sheet.merge_cells("A2:N2")

    cards = (
        "Total Sales",
        "Total Profit",
        "Total Orders",
        "Units Sold",
        "Total Customers",
        "Average Order Value",
        "Profit Margin",
    )
    kpi_values = kpis.set_index("Metric")["Value"]
    card_width = 2
    for index, metric in enumerate(cards):
        first_column = index * card_width + 1
        last_column = first_column + card_width - 1
        for row in range(4, 8):
            sheet.merge_cells(
                start_row=row,
                start_column=first_column,
                end_row=row,
                end_column=last_column,
            )
            sheet.cell(row=row, column=first_column).fill = PatternFill(
                "solid", fgColor="EAF1F5"
            )

        label = sheet.cell(row=4, column=first_column, value=metric.upper())
        label.font = Font(bold=True, color="526779", size=9)
        label.alignment = Alignment(vertical="center")
        value = sheet.cell(
            row=5,
            column=first_column,
            value=_excel_value(kpi_values.get(metric, pd.NA)),
        )
        value.font = Font(bold=True, color="187B86", size=18)
        value.alignment = Alignment(vertical="center")
        if metric == "Profit Margin":
            value.number_format = "0.0%;[Red]-0.0%;-"
        elif metric in {"Total Orders", "Units Sold", "Total Customers"}:
            value.number_format = "#,##0;[Red]-#,##0;-"
        else:
            value.number_format = "#,##0.00;[Red]-#,##0.00;-"

        definition = kpis.loc[kpis["Metric"].eq(metric), "Definition"].iloc[0]
        note = sheet.cell(row=7, column=first_column, value=str(definition))
        note.font = Font(color="526779", size=8)
        note.alignment = Alignment(vertical="top", wrap_text=True)

    quality = dict(zip(quality_report["Metric"], quality_report["Value"]))
    sheet.row_dimensions[7].height = 36
    sheet.merge_cells("A8:N8")
    quality_banner = sheet["A8"]
    quality_banner.value = (
        "DATA QUALITY  |  "
        f"{int(quality.get('Duplicate rows removed', 0)):,} duplicate rows removed   "
        f"•  {int(quality.get('Rows with missing or invalid order dates', 0)):,} invalid/missing dates   "
        f"•  {int(quality.get('Rows with negative sales', 0)):,} negative sales   "
        f"•  {int(quality.get('Rows with zero or negative quantity', 0)):,} nonpositive quantities"
    )
    quality_banner.font = Font(bold=True, color="8A5314", size=10)
    quality_banner.fill = PatternFill("solid", fgColor="FFF2D9")
    quality_banner.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[8].height = 29

    period_sections = (
        ("monthly", "MONTHLY SALES & GROWTH", "DashMonthly"),
        ("quarterly", "QUARTERLY SALES & GROWTH", "DashQuarterly"),
        ("yearly", "ANNUAL SALES & GROWTH", "DashYearly"),
        ("categories", "SALES BY CATEGORY", "DashCategories"),
        ("regions", "SALES BY REGION", "DashRegions"),
        ("products", "SALES BY PRODUCT", "DashProducts"),
        ("customers", "CUSTOMER PERFORMANCE", "DashCustomers"),
        ("quality", "DATA QUALITY DETAILS", "DashQuality"),
    )
    table_frames = {**analysis_tables, "quality": quality_report}
    table_locations: dict[str, tuple[int, int]] = {}
    section_ends: dict[str, int] = {}
    next_table_row = 78
    for key, title, table_name in period_sections:
        start_row = next_table_row
        start_column = 1
        table_locations[key] = (start_row + 1, start_column)
        section_ends[key] = _write_dashboard_table(
            sheet,
            title,
            table_frames[key],
            start_row,
            start_column,
            table_name,
        )
        next_table_row = section_ends[key] + 3

    _add_dashboard_insights(sheet, insights, 60)

    for table_key, title, anchor, chart_type in (
        ("monthly", "Monthly sales", "A10", "line"),
        ("categories", "Sales by category", "H10", "bar"),
        ("regions", "Sales by region", "A26", "bar"),
        ("products", "Top 10 products by sales", "H26", "bar"),
        ("customers", "Top 10 customers by sales", "A42", "bar"),
    ):
        frame = analysis_tables[table_key]
        if frame.empty or not frame["Sales"].notna().any():
            continue
        chart = LineChart() if chart_type == "line" else BarChart()
        chart.title = title
        chart.y_axis.title = "Sales (source units)"
        chart.style = 10
        chart.height = 6.2
        chart.width = 11.8
        start_row, start_column = table_locations[table_key]
        point_count = (
            min(len(frame), 10)
            if table_key in {"products", "customers"}
            else len(frame)
        )
        column_by_name = {
            column: start_column + index
            for index, column in enumerate(frame.columns)
        }
        sales_column = column_by_name["Sales"]
        dimension_column = column_by_name[frame.columns[0]]
        data = Reference(
            sheet,
            min_col=sales_column,
            min_row=start_row,
            max_row=start_row + point_count,
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(
            Reference(
                sheet,
                min_col=dimension_column,
                min_row=start_row + 1,
                max_row=start_row + point_count,
            )
        )
        if chart_type == "bar":
            chart.type = "bar"
            chart.legend = None
        else:
            chart.x_axis.title = "Month"
            chart.x_axis.number_format = "mmm yy"
        sheet.add_chart(chart, anchor)

    summary_end = max(section_ends.values(), default=next_table_row)
    cleaned_start = summary_end + 3
    _write_dashboard_table(
        sheet,
        "CLEANED SALES DETAIL  |  Use filters to find individual transactions.",
        sales,
        cleaned_start,
        1,
        "DashCleanedSales",
    )

    for column in range(1, 15):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = max(
            sheet.column_dimensions[letter].width or 0, 11
        )
    sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 0, 18)
    for column in range(1, 15):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = min(
            25, sheet.column_dimensions[letter].width
        )
    sheet.sheet_view.selection[0].activeCell = "A1"
    sheet.sheet_view.selection[0].sqref = "A1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:8"


def _write_excel_workbook(
    path: Path,
    sales: pd.DataFrame,
    quality_report: pd.DataFrame,
    kpis: pd.DataFrame,
    tables: dict[str, pd.DataFrame],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dashboard"
    insights = build_insights(sales, quality_report, tables, kpis)
    _add_dashboard_sheet(sheet, kpis, tables, quality_report, sales, insights)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _insight_summary(
    insights: str, charts: dict[str, Path], output_dir: Path
) -> None:
    (output_dir / "business_insights.md").write_text(insights, encoding="utf-8")
    for name, path in charts.items():
        if not path.is_file():
            raise FileNotFoundError(f"Expected chart was not created: {path}")


def generate_outputs(
    sales: pd.DataFrame,
    quality_report: pd.DataFrame,
    output_dir: str | Path = DEFAULT_OUTPUT_DIR,
) -> dict[str, Path]:
    """Write Power BI-ready data, interactive-workbook data, charts, and insights."""
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    tables = build_analysis_tables(sales)
    kpis = build_kpi_summary(sales, quality_report)

    csv_tables = {
        "cleaned_sales": sales,
        "kpi_summary": kpis,
        "data_quality_report": quality_report,
        "sales_by_month": tables["monthly"],
        "sales_by_quarter": tables["quarterly"],
        "sales_by_year": tables["yearly"],
        "sales_by_product": tables["products"],
        "sales_by_category": tables["categories"],
        "sales_by_region": tables["regions"],
        "sales_by_customer": tables["customers"],
    }
    outputs: dict[str, Path] = {}
    for name, frame in csv_tables.items():
        path = directory / f"{name}.csv"
        frame.to_csv(path, index=False, encoding="utf-8-sig", date_format="%Y-%m-%d")
        outputs[name] = path

    charts = _write_visualizations(tables, directory / "charts")
    outputs["charts"] = directory / "charts"
    outputs["business_insights"] = directory / "business_insights.md"
    _insight_summary(
        build_insights(sales, quality_report, tables, kpis), charts, directory
    )

    workbook_path = directory / "Sales_Analysis.xlsx"
    _write_excel_workbook(workbook_path, sales, quality_report, kpis, tables)
    outputs["excel_analysis"] = workbook_path
    return outputs


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description="Clean, analyze, and export the internship sales analytics project."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT_FILE,
        help="Source .xlsx or .csv sales dataset.",
    )
    parser.add_argument(
        "--sheet-name",
        default="Raw_Sales_Data",
        help="Worksheet to read when the source is an Excel workbook.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="Folder in which to save cleaned data, reports, charts, and the Excel workbook.",
    )
    args = parser.parse_args(argv)

    sales, quality_report = load_and_clean_data(args.input, args.sheet_name)
    outputs = generate_outputs(sales, quality_report, output_dir=args.output_dir)
    summary = build_kpi_summary(sales, quality_report)

    print(f"Source: {args.input.resolve()}")
    print(f"Cleaned rows: {len(sales):,}")
    print("\nKey performance indicators")
    for row in summary.itertuples(index=False):
        value = (
            f"{float(row[1]):.2%}"
            if row[0] == "Profit Margin" and not pd.isna(row[1])
            else _format_number(row[1])
        )
        print(f"  {row[0]}: {value}")
    print("\nProject outputs")
    for name, path in outputs.items():
        print(f"  {name}: {path.resolve()}")


if __name__ == "__main__":
    main()
